"""Parsing, column grouping and CSV export for workflow transcripts.

This mirrors the logic in the browser dashboard (../index.html) so both
versions turn the same transcript into the same handoff map.
"""
from __future__ import annotations

import io
import re
import uuid
from dataclasses import dataclass, field
from pathlib import Path

# ---------- Color key ----------
TYPE_GROUPS: list[tuple[str, list[tuple[str, str]]]] = [
    ("Reviewing", [("Reviewing AI Reasoning", "FFC5B8"), ("Reviewing Proposed Courses", "FF7A5C"),
                   ("Reviewing Proposed Calendar", "FF542E"), ("Reviewing Plan Revisions", "FF2F00"),
                   ("Reviewing AI Self-Check Results", "D12600"), ("Assess Available Resources", "A31E00")]),
    ("Consulting", [("Consulting Degree Progress Record", "FFFDB8"), ("Consulting Course Sequence Guide", "FFFA5C")]),
    ("AI work", [("Black Box", "FFB8FB")]),
    ("Cross-checking", [("Cross-Checking Prerequisites", "D8FFB8"), ("Cross-Checking Course Availability", "A5FF5C"),
                        ("Cross-Checking Credit Load", "8CFF2E"), ("Cross-Checking Against AI Reasoning", "73FF00"),
                        ("Cross-Checking Against Degree Progress", "5ED100"),
                        ("Cross-Checking Against Recommended Sequence", "357500")]),
    ("Other human tasks", [("Pause", "C4C4C4"), ("Drafting Prompt", "525252"), ("Approving", "FFC28A"),
                           ("Submitting Final Plan", "FF7B00"), ("Orienting to Next Step", "00FFEA")]),
    ("Handoffs", [("General Request", "9FABEA"), ("Specific Request", "7486E0"), ("Direction to Check", "546AD9"),
                  ("Direction to Proceed", "2E49D1"), ("Intermediate Product", "263CAB"),
                  ("Complete Product", "152260")]),
    # Higher-level task types. A name that is also a task type above (Black Box, Submitting Final Plan)
    # keeps that color, so it is not repeated here.
    ("Macro tasks", [("Reviewing and refining course recommendations", "ADFFD5"),
                     ("Validating the course plan", "ADCEFF"), ("Evaluating the proposed schedule", "FFF0AD"),
                     ("Evaluating the proposed calendar", "FFF0AD"),  # same macro task, as the sheets also write it
                     ("Decision Making", "FFADBC"), ("Finalizing the course plan", "DBADFF")]),
    # Micro tasks: neon colors, none used above. Related actions share a hue family.
    ("Micro tasks · AI and system", [("SESSION_START", "FF00DD"), ("AI_WORK_NOTICE", "E600B8"),
                                     ("AI_CONTENT_OUTPUT", "FF40B2"), ("ERROR_POPUP_DISPLAYED", "E6003D")]),
    ("Micro tasks · navigation", [("TAB_SWITCH", "0000FF"), ("PANEL_EXPAND", "0044FF"), ("PANEL_COLLAPSE", "4073FF"),
                                  ("OPEN_RESOURCE", "0099FF"), ("CLOSE_RESOURCE", "00A8E6")]),
    ("Micro tasks · expanding", [("EXPAND_PLAN_DIFF", "8800FF"), ("EXPAND_AI_REASONING", "A640FF"),
                                 ("EXPAND_SCHEDULE_CALENDAR", "DD00FF")]),
    ("Micro tasks · reading", [("SCROLL_AI_OUTPUT", "00DDFF"), ("SCROLL_REFERENCE_DOC", "40FFBF"),
                               ("HOVER_AI_OUTPUT", "00E68A"), ("HOVER_REFERENCE_DOC", "00FF88"),
                               ("CLICK_AI_OUTPUT_ITEM", "00FF66")]),
    ("Micro tasks · cursor", [("CURSOR_TRANSIT", "00E61F"), ("CURSOR_IDLE", "A8E600")]),
    ("Micro tasks · prompting", [("TYPE_PROMPT", "FFDD00"), ("EDIT_PROMPT_TEXT", "DDFF00"), ("SEND_PROMPT", "FFBB00")]),
    ("Micro tasks · steps", [("ADVANCE_STEP", "FFA640"), ("RETURN_STEP", "FF8040"), ("SUBMIT_FINAL", "FF0077")]),
]
# Views of a workflow: micro tasks, task types, macro (higher-level) tasks, and all three side by side
LEVELS = ("micro", "task", "high", "all")
LEVEL_NAME = {"micro": "Micro", "task": "Task", "high": "Macro"}
UNKNOWN_HEX = "#FFFFFF"
HIGH_HEADER = r"macro|high|parent"  # header of the macro (higher-level) task type column
MICRO_HEADER = r"micro"  # header of the micro task column
KIND_LABEL = {"human": "Human", "ai": "AI", "handoff": "Handoff"}


def norm(s) -> str:
    s = re.sub(r"\s+", " ", str("" if s is None else s).strip()).lower()
    return re.sub("[‐-―]", "-", s)


COLORS: dict[str, str] = {}
CANON: dict[str, str] = {}
for _, _items in TYPE_GROUPS:
    for _name, _hex in _items:
        COLORS[norm(_name)] = "#" + _hex
        CANON[norm(_name)] = _name


# ---------- Data model ----------
@dataclass
class Cell:
    name: str
    hex: str
    known: bool
    tasks: list[str] = field(default_factory=list)  # the level below it covers (macro: task types; task: micro tasks)
    group: str = ""  # the level above it (micro: task type; task: macro task)
    rows: int = 1  # how many micro task rows it spans (used by the all-levels view)
    parent: int = -1  # index of the square above it in the same column (micro -> task, task -> macro); -1: none


@dataclass
class Column:
    kind: str  # "human" | "ai" | "handoff"
    step: int
    label: str
    cells: list = field(default_factory=list)  # Cells; all-levels view also None (covered by a tall square above), "" (empty)
    high_cells: list[Cell] = field(default_factory=list)  # same column, one square per macro task
    micro_cells: list[Cell] = field(default_factory=list)  # same column, one square per micro task
    sub: str = ""  # all-levels view: which level this sub-column shows ("micro", "task", "high")


@dataclass
class Step:
    label: str
    start: int
    end: int


@dataclass
class Workflow:
    name: str
    category: str  # "tight" | "loose"
    columns: list[Column]
    steps: list[Step]
    warnings: list[str]
    counters: dict[str, int]
    id: str = field(default_factory=lambda: "w" + uuid.uuid4().hex[:7])
    has_high: bool = False  # the file has a macro (higher-level) task type column
    has_micro: bool = False  # the file has a micro task column

    @property
    def rows_max(self) -> int:
        return max((len(c.cells) for c in self.columns), default=0)

    def step_for(self, col_index: int) -> Step | None:
        return next((s for s in self.steps if s.start <= col_index <= s.end), None)

    @property
    def levels(self) -> list[str]:
        """The single-level views this file has, finest first."""
        return [lv for lv in ("micro", "task", "high") if lv == "task" or getattr(self, "has_" + lv)]

    def _copy(self, columns: list[Column], steps: list[Step]) -> "Workflow":
        return Workflow(self.name, self.category, columns, steps, self.warnings, self.counters, self.id,
                        self.has_high, self.has_micro)

    def at_level(self, level: str) -> "Workflow":
        """This workflow with its squares showing micro tasks, task types, macro tasks or all three ("all")."""
        if level == "all":
            return self._all_levels()
        if level == "task" or level not in self.levels:
            return self
        attr = "high_cells" if level == "high" else "micro_cells"
        columns = [Column(c.kind, c.step, c.label, getattr(c, attr), c.high_cells, c.micro_cells)
                   for c in self.columns]
        return self._copy(columns, self.steps)

    def _all_levels(self) -> "Workflow":
        """Each human and AI column split into one sub-column per level (micro, task, macro), one row per micro
        task, with task and macro squares as tall squares spanning the rows they cover."""
        levels = self.levels
        columns, first, last = [], [], []
        for c in self.columns:
            first.append(len(columns))
            if c.kind == "handoff":
                columns.append(Column(c.kind, c.step, c.label, list(c.cells), sub=""))
                last.append(len(columns) - 1)
                continue
            # Each row is a micro task; beside it, its task type and macro task ("" when it has none),
            # drawn once at the first row they cover (None on the rows below)
            tasks = [m.parent for m in c.micro_cells]
            highs = [c.cells[t].parent if t >= 0 else -1 for t in tasks]
            for lv in levels:
                if lv == "micro":
                    cells: list = list(c.micro_cells)
                else:
                    idx, src = (tasks, c.cells) if lv == "task" else (highs, c.high_cells)
                    cells = ["" if i < 0 else None if r and idx[r - 1] == i else src[i] for r, i in enumerate(idx)]
                columns.append(Column(c.kind, c.step, c.label, cells, sub=lv))
            last.append(len(columns) - 1)
        steps = [Step(s.label, first[s.start], last[s.end]) for s in self.steps]
        return self._copy(columns, steps)

    def to_dict(self) -> dict:
        """JSON shape used by the dashboard page (same fields as the browser version)."""
        cells = lambda xs: [{**vars(x), "tasks": list(x.tasks)} for x in xs]
        return {"id": self.id, "name": self.name, "category": self.category, "hasHigh": self.has_high,
                "hasMicro": self.has_micro,
                "columns": [{"kind": c.kind, "step": c.step, "label": c.label, "cells": cells(c.cells),
                             "highCells": cells(c.high_cells), "microCells": cells(c.micro_cells)}
                            for c in self.columns],
                "steps": [vars(s).copy() for s in self.steps], "warnings": list(self.warnings),
                "counters": dict(self.counters), "rowsMax": self.rows_max,
                "highRowsMax": self.at_level("high").rows_max, "microRowsMax": self.at_level("micro").rows_max}

    @classmethod
    def from_dict(cls, d: dict) -> "Workflow":
        cells = lambda xs: [Cell(str(x["name"]), str(x["hex"]), bool(x["known"]),
                                 [str(t) for t in x.get("tasks") or []], str(x.get("group") or ""),
                                 max(1, int(x.get("rows") or 1)), int(x.get("parent", -1))) for x in xs]
        columns = [Column(c["kind"], int(c["step"]), c["label"], cells(c["cells"]),
                          cells(c.get("highCells") or c["cells"]), cells(c.get("microCells") or c["cells"]))
                   for c in d["columns"]]
        steps = [Step(s["label"], int(s["start"]), int(s["end"])) for s in d["steps"]]
        return cls(d["name"], d["category"], columns, steps, list(d.get("warnings", [])),
                   dict(d["counters"]), d["id"], bool(d.get("hasHigh")), bool(d.get("hasMicro")))


# ---------- Parsing ----------
def parse_csv(text: str) -> list[list[str]]:
    """Read CSV text, guessing the delimiter (comma, semicolon or tab) from the first line."""
    text = text.lstrip("﻿")
    first = re.split(r"\r?\n", text, maxsplit=1)[0]
    counts = {d: first.count(d) for d in (",", ";", "\t")}
    delim = max(counts, key=lambda d: counts[d])
    rows, row, cell, quoted = [], [], "", False
    i, n = 0, len(text)
    while i < n:
        c = text[i]
        if quoted:
            if c == '"':
                if i + 1 < n and text[i + 1] == '"':
                    cell += '"'
                    i += 1
                else:
                    quoted = False
            else:
                cell += c
        elif c == '"':
            quoted = True
        elif c == delim:
            row.append(cell)
            cell = ""
        elif c in "\r\n":
            if c == "\r" and i + 1 < n and text[i + 1] == "\n":
                i += 1
            row.append(cell)
            rows.append(row)
            row, cell = [], ""
        else:
            cell += c
        i += 1
    if cell != "" or row:
        row.append(cell)
        rows.append(row)
    return rows


def parse_xlsx(data: bytes) -> list[list[str]]:
    from openpyxl import load_workbook

    ws = load_workbook(io.BytesIO(data), read_only=True, data_only=True).worksheets[0]
    return [["" if v is None else str(v) for v in r] for r in ws.iter_rows(values_only=True)]


def read_rows(filename: str, data: bytes) -> list[list[str]]:
    if re.search(r"\.xlsx?$", filename, re.I):
        rows = parse_xlsx(data)
    else:
        rows = parse_csv(data.decode("utf-8-sig", errors="replace"))
    if not rows or all(not str(v).strip() for r in rows for v in r):
        raise ValueError("the file is empty")
    return rows


def base_name(filename: str) -> str:
    return re.sub(r"\.(csv|xlsx|xls)$", "", Path(filename).name, flags=re.I)


def build_workflow(rows: list[list[str]], name: str) -> Workflow:
    warnings: list[str] = []
    rows = [["" if v is None else str(v).strip() for v in r] for r in rows]
    header = rows[0] if rows else []
    get = lambda r, i: r[i] if 0 <= i < len(r) else ""

    # Optional columns: Micro Task (left of Task Type) and Macro Task (between Task Type and Handoff Type)
    find = lambda pattern, skip=(): next((i for i, h in enumerate(header)
                                          if i and i not in skip and re.search(pattern, norm(h))), -1)
    micro_col = find(MICRO_HEADER)
    high_col = find(HIGH_HEADER, (micro_col,))
    extra = (micro_col, high_col)

    def find_col(pattern: str, fallback: int) -> int:
        i = find(pattern, extra)
        return i if i >= 0 else fallback

    task_col, hand_col, agent_col = find_col("task", 1), find_col("hand", 2), find_col("agent", 3)
    if find("task", extra) < 0:
        warnings.append("No “Task Type” header found, so columns B, C and D were read as "
                        "Task Type, Handoff Type and Agent.")

    has_step_labels = any(get(r, 0) != "" for r in rows[1:])
    cat_raw = norm(get(header, 0))
    if cat_raw.startswith("tight"):
        category = "tight"
    elif cat_raw.startswith("loose"):
        category = "loose"
    else:
        category = "tight" if has_step_labels else "loose"
        warnings.append(f"The first cell says “{get(header, 0) or '(blank)'}”. Treated as {category} because "
                        f"the file {'has' if has_step_labels else 'has no'} step labels.")

    # Flatten rows into a sequence of entries. A task entry is
    # [kind, task type, step, macro task or "", micro tasks, orphan]; a handoff has no macro or micro tasks.
    # An orphan is a micro task with no task type above it (e.g. SESSION_START): it shows as a micro task only.
    # A row with a task type starts a task; a row with only a micro task adds it to the task above.
    entries: list[list] = []
    step_idx, step_labels = -1, []
    last_agent, inherited, orphans, seen_task = None, 0, [], False
    for r in rows[1:]:
        new_step = category == "tight" and get(r, 0)
        if new_step:
            step_labels.append(get(r, 0))
            step_idx = len(step_labels) - 1
        task, hand, agent_raw = get(r, task_col), get(r, hand_col), norm(get(r, agent_col))
        micro = get(r, micro_col) if micro_col >= 0 else ""
        last = entries[-1] if entries else None
        if micro and not task and last and last[0] != "handoff" and not new_step:
            last[4].append(micro)
        elif task or micro:
            if re.search(r"\bai\b|^ai|artificial", agent_raw):
                kind = "ai"
            elif re.search("human|person|user", agent_raw):
                kind = "human"
            else:
                kind = last_agent or "human"
                inherited += 1
            last_agent = kind
            if not task and seen_task:  # micro tasks before the first task type (session events) need no note
                orphans.append(micro)
            seen_task = seen_task or bool(task)
            entries.append([kind, task or micro, step_idx, get(r, high_col) if high_col >= 0 else "",
                            [micro] if micro else [], not task])
        if hand:
            entries.append(["handoff", hand, step_idx, "", [], False])
    if inherited:
        one = inherited == 1
        warnings.append(f"{inherited} task {'row has' if one else 'rows have'} no agent, so "
                        f"{'it was' if one else 'they were'} kept with the agent of the task above.")
    if category == "tight" and not step_labels:
        warnings.append("Marked Tight but no step labels were found in the first column.")
    if category == "tight" and any(e[2] < 0 for e in entries):
        warnings.append("Some rows come before the first step label and sit outside any step.")
    if orphans:
        one = len(orphans) == 1
        warnings.append(f"{', '.join(dict.fromkeys(orphans))} {'has' if one else 'have'} no task type above "
                        f"{'it' if one else 'them'} in {'its' if one else 'their'} column, so "
                        f"{'it shows' if one else 'they show'} as {'a micro task' if one else 'micro tasks'} only.")

    # Group consecutive entries of the same kind within the same step
    counters = {"human": 0, "ai": 0, "handoff": 0}
    columns: list[Column] = []
    unknown: list[str] = []
    high_name, high_kind, untyped = "", None, 0

    def cell(entry_name: str) -> Cell:
        key = norm(entry_name)
        hex_ = COLORS.get(key)
        if not hex_ and entry_name not in unknown:
            unknown.append(entry_name)
        return Cell(CANON.get(key, entry_name), hex_ or UNKNOWN_HEX, bool(hex_))

    for kind, entry_name, step, high, micros, orphan in entries:
        last = columns[-1] if columns else None
        new_col = not last or last.kind != kind or last.step != step
        if new_col:
            counters[kind] += 1
            columns.append(Column(kind, step, f"{KIND_LABEL[kind]} #{counters[kind]}"))
        col = columns[-1]
        if orphan:  # micro tasks only: no task type or macro task square
            for m in micros:
                col.micro_cells.append(cell(m))
            continue
        x = cell(entry_name)
        col.cells.append(x)
        if kind == "handoff":
            col.high_cells.append(x)
            col.micro_cells.append(x)
            continue
        # Micro tasks: one square each; a task with none keeps its task type square
        x.tasks, x.rows = [m for m in micros], max(1, len(micros))
        for m in micros or [None]:
            mx = cell(m) if m else Cell(x.name, x.hex, x.known)
            mx.group, mx.parent = x.name, len(col.cells) - 1
            col.micro_cells.append(mx)
        # A macro task covers the task rows below it until the next one is named.
        # Blank rows continue it while the agent stays the same (a new column starts a new square
        # with the same type); a blank row after another agent's task has none of its own.
        if high or high_kind != kind:
            high_name = high
        high_kind = kind
        if high_col >= 0 and not high_name:
            untyped += 1
        if high or new_col or not high_name:
            hx = cell(high_name or entry_name)
            hx.rows = 0
            col.high_cells.append(hx)
        hx = col.high_cells[-1]
        hx.tasks.append(x.name)
        hx.rows += x.rows
        x.group, x.parent = hx.name, len(col.high_cells) - 1
    if unknown:
        warnings.append(f"No color is set for: {', '.join(unknown)}. These show as striped squares "
                        f"and export as {UNKNOWN_HEX}.")
    if untyped:
        one = untyped == 1
        warnings.append(f"{untyped} task {'row has' if one else 'rows have'} no macro task (blank, "
                        f"with no task by the same agent above), so {'it uses its' if one else 'they use their'} "
                        f"own task type in the macro view.")

    steps: list[Step] = []
    if category == "tight":
        for i, label in enumerate(step_labels):
            idx = [j for j, c in enumerate(columns) if c.step == i]
            if idx:
                steps.append(Step(label, idx[0], idx[-1]))
            else:
                warnings.append(f"“{label}” has no rows under it.")
    return Workflow(name, category, columns, steps, warnings, counters, has_high=high_col >= 0,
                    has_micro=micro_col >= 0)


def load_workflow(filename: str, data: bytes, name: str | None = None) -> Workflow:
    return build_workflow(read_rows(filename, data), name or base_name(filename))


# ---------- Tabs ----------
@dataclass
class Tab:
    id: str
    title: str
    kind: str  # "tight" | "loose" | "all"
    files: list[Workflow]
    level: str = "task"  # one of LEVELS

    @property
    def levels(self) -> list[str]:
        """Views worth exporting: every level any file has, then all levels side by side when there are several."""
        have = {lv for f in self.files for lv in f.levels}
        return [lv for lv in LEVELS if lv in have or (lv == "all" and len(have) > 1)]

    def at_level(self, level: str) -> "Tab":
        return Tab(self.id, self.title, self.kind, [f.at_level(level) for f in self.files], level)


def ordered(workflows: list[Workflow]) -> list[Workflow]:
    return [w for w in workflows if w.category == "tight"] + [w for w in workflows if w.category == "loose"]


def make_tabs(workflows: list[Workflow]) -> list[Tab]:
    tabs = [Tab(w.id, w.name, w.category, [w]) for w in workflows]
    if len(workflows) > 1:
        tabs.append(Tab("all", "All workflows", "all", ordered(workflows)))
    return tabs


def groups_of(tab: Tab) -> list[tuple[str, list[Workflow]]]:
    groups = [(cat, [f for f in tab.files if f.category == cat]) for cat in ("tight", "loose")]
    return [g for g in groups if g[1]]


def legend_of(files: list[Workflow]) -> list[Cell]:
    seen: dict[str, Cell] = {}
    for f in files:
        for c in f.columns:
            for x in c.cells:
                if x:
                    seen.setdefault(x.name, x)
    return list(seen.values())


def slug(s: str) -> str:
    s = re.sub(r"\s+", "-", re.sub(r"[^\w\- ]+", "", s).strip())
    return s or "workflow"


def file_base(tab: Tab) -> str:
    base = "all-workflows" if tab.id == "all" else slug(tab.title)
    return base + LEVEL_SUFFIX[tab.level]


LEVEL_SUFFIX = {"task": "", "micro": "-micro", "high": "-macro", "all": "-all-levels"}


def tab_levels(tab: Tab) -> list[Tab]:
    """The tab once per version to export: task types first, then micro tasks, macro tasks and all levels
    side by side, for whichever the files have."""
    order = ["task"] + [lv for lv in tab.levels if lv != "task"]
    return [tab.at_level(lv) for lv in order]


# ---------- CSV export ----------
def csv_cell(v) -> str:
    v = "" if v is None else str(v)
    return '"' + v.replace('"', '""') + '"' if re.search(r'[",\n\r]', v) else v


def step_row(f: Workflow) -> list[str]:
    out = [""] * len(f.columns)
    for s in f.steps:
        if s.start == s.end:
            out[s.start] = f"START/END: {s.label}"
        else:
            out[s.start] = f"START: {s.label}"
            out[s.end] = f"END: {s.label}"
    return out


def _pad(row: list[str], width: int) -> list[str]:
    return row + [""] * (width - len(row))


def block_matrix(f: Workflow, with_title: bool) -> list[list[str]]:
    """Rows of equal width for one workflow."""
    w = max(1, len(f.columns))
    rows = []
    if with_title:
        rows.append(_pad([f"{f.name} ({f.category})"], w))
    if f.category == "tight" or with_title:
        rows.append(_pad(step_row(f) if f.category == "tight" else [], w))
    rows.append(_pad([f"{c.label} · {LEVEL_NAME[c.sub]}" if c.sub else c.label for c in f.columns], w))
    # A tall square (all-levels view) repeats its color on every row it spans
    held: list[str] = [""] * len(f.columns)
    for r in range(f.rows_max):
        out = []
        for i, c in enumerate(f.columns):
            x = c.cells[r] if r < len(c.cells) else ""
            held[i] = x.hex if x else (held[i] if x is None else "")
            out.append(held[i])
        rows.append(_pad(out, w))
    return rows


def tab_csv(tab: Tab) -> str:
    if tab.id != "all":
        grid = block_matrix(tab.files[0], False)
    else:
        # Workflows stacked top to bottom: tight group first, then loose
        w = max([1] + [len(f.columns) for f in tab.files])
        grid = []
        for gi, (cat, files) in enumerate(groups_of(tab)):
            if gi:
                grid += [[""] * w, [""] * w]
            grid.append(_pad([cat.upper()], w))
            for fi, f in enumerate(files):
                if fi:
                    grid.append([""] * w)
                grid += [_pad(r, w) for r in block_matrix(f, True)]
    lines = [",".join(csv_cell(v) for v in r) for r in grid]
    lines += ["", "Legend,Hex"] + [f"{csv_cell(x.name)},{x.hex}" for x in legend_of(tab.files)]
    return "\r\n".join(lines) + "\r\n"
